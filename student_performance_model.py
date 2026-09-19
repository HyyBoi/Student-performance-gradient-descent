"""
Student performance model: multivariable linear regression trained with
batch gradient descent, written in plain Python (no NumPy) so every step of
the calculus is visible.

Model:   F(x1, x2, x3, x4) = w1*x1 + w2*x2 + w3*x3 + w4*x4 + b
Cost:    J = (1/m) * sum((F(x_i) - y_i)^2)
Gradient: dJ/dw_j = (2/m) * sum((F(x_i) - y_i) * x_ij),  dJ/db = (2/m) * sum(F(x_i) - y_i)

Usage:
    python student_performance_model.py --demo
    python student_performance_model.py --train Dataset_1.xlsx --test Dataset_2.xlsx
"""

import argparse
import math
import os
import random

FEATURE_NAMES = ["Assignments", "Quizzes", "Sessional-I", "Sessional-II"]
PASS_THRESHOLD = 50.0  # predicted marks out of 100 needed to count as PASS


# ---------------------------------------------------------------------------
# Model, cost and gradient
# ---------------------------------------------------------------------------

def predict(x, w, b):
    """F(x) = w . x + b"""
    result = b
    for i in range(len(w)):
        result += w[i] * x[i]
    return result


def compute_cost(X, Y, w, b):
    """Mean squared error J(w, b)."""
    m = len(Y)
    total_error = 0.0
    for i in range(m):
        error = predict(X[i], w, b) - Y[i]
        total_error += error ** 2
    return total_error / m


def compute_gradients(X, Y, w, b):
    """Return (dJ/dw as a list, dJ/db) at the current parameters."""
    m = len(Y)
    n = len(w)
    dw = [0.0] * n
    db = 0.0

    for i in range(m):
        error = predict(X[i], w, b) - Y[i]
        for j in range(n):
            dw[j] += error * X[i][j]
        db += error

    dw = [(2 / m) * g for g in dw]
    db = (2 / m) * db
    return dw, db


def gradient_descent(X, Y, learning_rate, num_iterations):
    """Steepest descent: repeatedly step against the gradient of J."""
    w = [0.0] * len(FEATURE_NAMES)
    b = 0.0

    print(f"  Learning Rate = {learning_rate}")
    print(f"  Iterations    = {num_iterations}")
    print()

    for iteration in range(num_iterations):
        dw, db = compute_gradients(X, Y, w, b)

        for j in range(len(w)):
            w[j] = w[j] - learning_rate * dw[j]
        b = b - learning_rate * db

        if (iteration + 1) % 1000 == 0 or iteration == 0:
            cost = compute_cost(X, Y, w, b)
            print(f"  Iteration {iteration + 1:>6} | Cost (MSE): {cost:.4f}")

    return w, b


# ---------------------------------------------------------------------------
# Data loading
# ---------------------------------------------------------------------------

def load_training_data(path, sheet="Fall 2025"):
    """Read the course-provided training workbook (column positions are specific to that sheet)."""
    import openpyxl  # imported here so --demo works without openpyxl installed

    ws = openpyxl.load_workbook(path)[sheet]
    rows = list(ws.iter_rows(values_only=True))

    X, Y = [], []
    for row in rows[7:]:  # first 7 rows are headers
        x1 = row[6]   # Assignments (weighted)
        x2 = row[19]  # Quizzes (weighted)
        x3 = row[22]  # Sessional-I (weighted)
        x4 = row[25]  # Sessional-II (weighted)
        y = row[32]   # Grand total

        if None in (x1, x2, x3, x4, y):
            continue  # skip students with incomplete training rows
        X.append([x1, x2, x3, x4])
        Y.append(y)
    return X, Y


def load_test_data(path, sheet="Spring 2026"):
    """Read the unseen workbook. Missing marks are treated as 0 and flagged."""
    import openpyxl

    ws = openpyxl.load_workbook(path)[sheet]
    rows = list(ws.iter_rows(values_only=True))

    students = []
    for row in rows[7:]:
        raw = [row[3], row[12], row[14], row[18]]
        has_missing = any(v is None for v in raw)
        features = [0.0 if v is None else v for v in raw]
        students.append((features, has_missing))
    return students


def make_demo_data(n, seed):
    """Synthetic students so the repo runs without the (private) course data."""
    rng = random.Random(seed)
    X, Y = [], []
    for _ in range(n):
        x = [rng.uniform(0, 8), rng.uniform(0, 10), rng.uniform(0, 15), rng.uniform(0, 15)]
        y = 2.5 * x[0] + 1.5 * x[1] + 2.0 * x[2] + 2.0 * x[3] + 1.0 + rng.gauss(0, 4)
        X.append(x)
        Y.append(max(0.0, min(100.0, y)))
    return X, Y


# ---------------------------------------------------------------------------
# Main workflow
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="Gradient-descent student performance model")
    parser.add_argument("--train", default="Dataset_1.xlsx", help="training workbook")
    parser.add_argument("--test", default="Dataset_2.xlsx", help="unseen workbook to predict on")
    parser.add_argument("--iterations", type=int, default=5000)
    parser.add_argument("--demo", action="store_true", help="use synthetic data instead of Excel files")
    args = parser.parse_args()

    # 1. Load data ----------------------------------------------------------
    if args.demo:
        X_train, Y_train = make_demo_data(60, seed=1)
        X_test, _ = make_demo_data(20, seed=2)
        test_students = [(x, False) for x in X_test]
        print("Running in DEMO mode with synthetic data.\n")
    else:
        for path in (args.train, args.test):
            if not os.path.exists(path):
                raise SystemExit(
                    f"Could not find '{path}'. The course datasets are not included in this repo.\n"
                    "Run with --demo to use synthetic data, or pass your own files with --train / --test."
                )
        X_train, Y_train = load_training_data(args.train)
        test_students = load_test_data(args.test)
    print(f"Students loaded for training: {len(X_train)}\n")

    # 2. Try several learning rates ------------------------------------------
    print("=" * 60)
    print("Gradient Descent with Different Learning Rates")
    print("=" * 60)
    print()

    best_w, best_b, best_cost, best_alpha = None, None, float("inf"), None

    for alpha in [0.001, 0.01, 0.1]:
        print(f"--- Training with alpha = {alpha} ---")
        w, b = gradient_descent(X_train, Y_train, alpha, args.iterations)
        cost = compute_cost(X_train, Y_train, w, b)
        print(f"  Final Cost: {cost:.4f}")
        print(f"  Weights: w1={w[0]:.4f}, w2={w[1]:.4f}, w3={w[2]:.4f}, w4={w[3]:.4f}")
        print(f"  Bias: b={b:.4f}\n")

        if math.isnan(cost) or math.isinf(cost):
            print("  *** DIVERGED: learning rate too large, weights exploded ***\n")
            continue

        if cost < best_cost:
            best_w, best_b, best_cost, best_alpha = w, b, cost, alpha

    print("=" * 60)
    print(f"Best learning rate: alpha = {best_alpha}")
    print(f"Best final cost (MSE): {best_cost:.4f}  (RMSE ~ {math.sqrt(best_cost):.2f} marks)")
    for name, weight in zip(FEATURE_NAMES, best_w):
        print(f"  {name:<14} weight = {weight:.4f}")
    print(f"  {'Bias':<14}        = {best_b:.4f}")
    print("=" * 60)
    print()

    # 3. Gradient at the solution --------------------------------------------
    print("Gradient vector at the trained weights (should be ~0 at a minimum):")
    dw, db = compute_gradients(X_train, Y_train, best_w, best_b)
    for name, g in zip(FEATURE_NAMES, dw):
        print(f"  dJ/dw ({name}) = {g:.6f}")
    print(f"  dJ/db (Bias) = {db:.6f}\n")

    top = max(range(len(best_w)), key=lambda j: abs(best_w[j]))
    print(f"Largest raw weight: {FEATURE_NAMES[top]} ({best_w[top]:.4f})")
    print("(Note: raw weights are only comparable if inputs are on similar scales.)\n")

    # 4. Predict on unseen data ------------------------------------------------
    print("=" * 60)
    print(f"Predictions on unseen data (pass mark = {PASS_THRESHOLD:.0f}/100)")
    print("=" * 60)
    print(f"{'Student':<12} {'As':>6} {'Qz':>6} {'S-I':>6} {'S-II':>6} | {'Predicted':>10} {'Result':>8}")
    print("-" * 65)

    num_pass = num_fail = 0
    for idx, (x, has_missing) in enumerate(test_students, start=1):
        predicted = max(0.0, min(100.0, predict(x, best_w, best_b)))  # clamp to 0-100
        result = "PASS" if predicted >= PASS_THRESHOLD else "FAIL"
        if result == "PASS":
            num_pass += 1
        else:
            num_fail += 1
        note = " (has missing data)" if has_missing else ""
        print(f"Student {idx:<4} {x[0]:>6.2f} {x[1]:>6.2f} {x[2]:>6.2f} {x[3]:>6.2f} | "
              f"{predicted:>10.2f} {result:>8}{note}")

    print("-" * 65)
    print(f"Total students predicted: {num_pass + num_fail}")
    print(f"  Predicted PASS: {num_pass}")
    print(f"  Predicted FAIL: {num_fail}")


if __name__ == "__main__":
    main()